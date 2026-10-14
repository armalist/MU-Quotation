# CATATAN: script ini dipakai sekali untuk membuat rate card dari 761 PH historis (butuh file hist.pkl hasil parsing quotation).
# Untuk memperbarui harga sehari-hari TIDAK perlu script ini: cukup Upload database di aplikasi.
# Bangun database rate card default dari hist.pkl -> JSON (untuk HTML), XLSX, CSV, template upload
import pandas as pd, json, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
W = sys.argv[1]
H = pd.read_pickle(W + '/q/hist.pkl')
CATS = [('FD','Floor Display - POSM'),('FSU','FSU - POSM'),('EG','End Gondola - POSM'),('WG','Wagon - POSM'),('CTA','CTA - POSM'),
        ('DP','Display - POSM'),('SB','Store Build & Refurbish'),('BE','Booth & Event'),('BS','Branding & Signage'),('RV','Revisual & Repostering'),
        ('IN','Installasi'),('DM','Dismantle'),('MT','Maintenance'),('MF','Maintenance & FSU - POSM'),('ME','Maintenance & End Gondola - POSM'),
        ('SV','Survey'),('OJ','Operasional & Jasa')]
H['Kategori'] = H.Kode.map(dict(CATS)); H['o'] = H.Kode.map({k:i for i,(k,_) in enumerate(CATS)})
H = H.sort_values(['o','Subkategori','nph','Nama'], ascending=[True,True,False,True]).reset_index(drop=True)
def clean(x):
    x = '' if x is None or (isinstance(x,float) and pd.isna(x)) else str(x).strip()
    return '' if x in ('-','nan','None') else x
rows, seen = [], {}
for d in H.to_dict('records'):
    nama = clean(d['Nama'])[:90]; uk = clean(d['UkuranRC']); sat = clean(d['Satuan'])
    k = (d['Kategori'], d['Subkategori'], nama.lower(), uk.lower(), sat.lower()); n = seen.get(k,0)+1; seen[k]=n
    if n > 1: nama = f'{nama} #{n}'
    tg = d['Tanggal'].strftime('%Y-%m-%d') if pd.notna(d['Tanggal']) else ''
    r = lambda v: int(round(float(v))) if pd.notna(v) else 0
    rows.append([d['Kategori'], d['Subkategori'], nama, uk, sat, r(d['mn']), r(d['md']), r(d['mx']), 0, int(d['nph']), r(d['unit_price']), tg, clean(d.get('Klien')), clean(d.get('No PH'))])
HDR = ['Kategori','Sub Kategori','Nama Item','Ukuran','Satuan','Harga Min','Harga Acuan','Harga Maks','Modal/HPP','Jumlah PH','Harga Terakhir','Tgl Terakhir','Klien Terakhir','No PH Terakhir']
asof = max(r[11] for r in rows)
json.dump({'rows':rows,'source':'Rate card historis 761 PH terverifikasi (2025–2026)','asof':asof}, open(W+'/app/db_default.json','w'), ensure_ascii=False, separators=(',',':'))
pd.DataFrame(rows, columns=HDR).to_csv(W+'/app/database/Ratecard_Database_MU.csv', index=False, encoding='utf-8-sig')

NAVY='FF1F3864'; thin=Side(style='thin',color='FFA6A6A6'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
WID=[24,32,60,18,9,14,14,14,13,10,14,12,34,20]
def sheet(ws, data, note):
    for j,h in enumerate(HDR,1):
        c=ws.cell(1,j,h); c.font=Font(name='Calibri',bold=True,color='FFFFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY if j!=9 else 'FFC55A11')
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
        ws.column_dimensions[c.column_letter].width=WID[j-1]
    ws.row_dimensions[1].height=30
    for i,r in enumerate(data,2):
        for j,v in enumerate(r,1):
            if j==12 and v: v=dt.datetime.strptime(v,'%Y-%m-%d')
            if v in ('',0) and j in (9,): v=None
            c=ws.cell(i,j,v if v!='' else None); c.font=Font(name='Calibri',size=10,bold=(j==7)); c.border=BOX
            if j in (6,7,8,9,11): c.number_format='#,##0'
            if j==12: c.number_format='d mmm yyyy'
    ws.freeze_panes='D2'; ws.auto_filter.ref=f'A1:N{len(data)+1}'
def guide(ws, extra):
    ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=95
    L=[('PETUNJUK DATABASE RATE CARD',''),('',''),
       ('Fungsi','File ini dibaca oleh Generator PH (HTML) lewat tombol "Upload database". Sheet yang dibaca: sheet pertama yang punya kolom Kategori & Nama Item.'),
       ('Wajib diisi','Kategori, Nama Item, dan minimal salah satu: Harga Acuan atau Modal/HPP.'),
       ('Kategori','Pakai 17 nama kategori Mata Utama (mis. "End Gondola - POSM", "Branding & Signage"). Nama baru tetap diterima dan akan muncul sebagai kategori baru.'),
       ('Sub Kategori','Boleh kosong (akan masuk "Lainnya"). Pakai nama yang sama persis agar item berkumpul di dropdown yang sama.'),
       ('Harga Min/Acuan/Maks','Harga JUAL historis ke klien, sebelum PPN, per satuan. Acuan = median. Kalau hanya isi satu harga, Min & Maks disamakan.'),
       ('Modal/HPP','Opsional (kolom oranye). Kalau diisi, harga di PH = Modal × (1 + margin). Isi dari price list purchasing/vendor.'),
       ('Kunci item','Kategori + Sub Kategori + Nama Item + Ukuran + Satuan. Mode "Update" menimpa item dengan kunci sama & menambah item baru. Mode "Ganti semua" mengganti seluruh database.'),
       ('Format angka','Angka biasa (1250000) atau teks Indonesia (1.250.000) sama-sama bisa. Tanggal: tanggal Excel atau YYYY-MM-DD.'),
       ('CSV','Boleh juga simpan sebagai CSV (pemisah koma atau titik koma), header sama.')]+extra
    for i,(a,b) in enumerate(L,1):
        ws.cell(i,1,a).font=Font(name='Calibri',bold=True,size=14 if i==1 else 11); c=ws.cell(i,2,b); c.font=Font(name='Calibri'); c.alignment=Alignment(wrap_text=True,vertical='top')
wb=Workbook(); ws=wb.active; ws.title='Rate Card'; sheet(ws, rows, ''); guide(wb.create_sheet('Petunjuk'), [('Sumber',f'{len(rows):,} item dari 761 PH terverifikasi folder Quotation 2025–2026. Data s/d {asof}.'.replace(',','.'))])
wb.save(W+'/app/database/Ratecard_Database_MU.xlsx')
sample=[['End Gondola - POSM','End Gondola','Contoh: Produksi End Gondola Multiplek 15mm','120 x 40 x 200 cm','unit',0,8500000,0,6200000,0,0,'','',''],
        ['Branding & Signage','Neon Box / Lightbox','Contoh: Neon Box PVC 8mm + LED','94 x 99 cm','unit',1800000,2000000,2300000,0,0,0,'','',''],
        ['Installasi','Instalasi','Contoh: Biaya Install & Transport Jabodetabek','','toko',600000,650000,750000,0,0,0,'','','']]
wb=Workbook(); ws=wb.active; ws.title='Rate Card'; sheet(ws, sample, ''); guide(wb.create_sheet('Petunjuk'), [('Template','Hapus 3 baris contoh, isi item Anda mulai baris 2. Jangan ubah nama kolom di baris 1.')])
wb.save(W+'/app/database/Template_Upload_Ratecard.xlsx')
print(len(rows), asof, sum(1 for r in rows if '#' in r[2]))
